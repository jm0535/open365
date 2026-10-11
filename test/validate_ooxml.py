#!/usr/bin/env python3
"""Independent OOXML check for the files written by ooxml.test.mjs.

Uses python-docx and openpyxl — not the suite's own ZIP reader — so a
serializer bug that both writes and reads the same mistake still fails CI.
"""
import os
import sys
from pathlib import Path

docx_path = Path('/tmp/t.docx')
xlsx_path = Path('/tmp/t.xlsx')

if not docx_path.exists() or not xlsx_path.exists():
    print('missing /tmp/t.docx or /tmp/t.xlsx — ooxml.test.mjs did not write them', file=sys.stderr)
    sys.exit(1)

try:
    import docx as docxmod
    import openpyxl
except ImportError:
    print('python-docx and openpyxl are required', file=sys.stderr)
    sys.exit(1 if os.environ.get('CI') else 0)

document = docxmod.Document(str(docx_path))
texts = [p.text for p in document.paragraphs]
if not any('Quarterly report' in t for t in texts):
    print('docx heading missing:', texts, file=sys.stderr)
    sys.exit(1)
styles = [p.style.name for p in document.paragraphs if p.style]
if not any('Heading' in (s or '') for s in styles):
    print('docx heading style missing:', styles, file=sys.stderr)
    sys.exit(1)
bold = False
for p in document.paragraphs:
    for run in p.runs:
        if run.bold and 'bold' in (run.text or ''):
            bold = True
if not bold:
    print('docx bold run missing', file=sys.stderr)
    sys.exit(1)

wb = openpyxl.load_workbook(str(xlsx_path))
ws = wb.active
formula = ws.cell(3, 4).value
if formula != '=C3-B3':
    print('xlsx formula mismatch:', formula, file=sys.stderr)
    sys.exit(1)
total = ws.cell(6, 2).value
if total != '=SUM(B3:B5)':
    print('xlsx SUM mismatch:', total, file=sys.stderr)
    sys.exit(1)
fmt = ws.cell(3, 4).number_format or ''
if '¥' in fmt or fmt in ('General', '0', ''):
    # currency style must not collapse to General
    if not any(ch in fmt for ch in ('$', '¤', '#', '0')):
        print('xlsx currency format missing:', fmt, file=sys.stderr)
        sys.exit(1)

print('OOXML python validators passed')
